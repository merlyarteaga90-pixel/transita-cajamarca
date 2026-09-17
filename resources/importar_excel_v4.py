"""
Importador del Excel v4 al MySQL.

Lee tbl_rutas y tbl_itinerarios del archivo versionado:
resources/listado_rutas_cajamarca_2024.xlsx
(se puede sobreescribir con RUTAS_EXCEL_PATH o pasando la ruta como argumento).

Ejecutar desde la raiz del proyecto:
    python resources/importar_excel_v4.py [ruta_excel_opcional]

Requiere:
    pip install openpyxl sqlalchemy pymysql
"""

import os
import re
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from dotenv import load_dotenv

import openpyxl
from sqlalchemy import create_engine, text
from sqlalchemy.exc import IntegrityError

load_dotenv()

PROYECTO = Path(__file__).resolve().parent.parent
EXCEL_PATH = os.getenv(
    "RUTAS_EXCEL_PATH",
    str(PROYECTO / "resources" / "listado_rutas_cajamarca_2024.xlsx"),
)

DB_USER = os.getenv("DB_USER", "root")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_HOST = os.getenv("DB_HOST", "localhost")
DB_PORT = os.getenv("DB_PORT", "3307")
DB_NAME = os.getenv("DB_NAME", "asistente_rutas")

DATABASE_URL = (
    f"mysql+pymysql://{DB_USER}:{DB_PASSWORD}"
    f"@{DB_HOST}:{DB_PORT}/{DB_NAME}?charset=utf8mb4"
)


# ============================================================
# NORMALIZACION
# ============================================================

def normalizar(texto: str) -> str:
    if not texto:
        return ""
    texto = str(texto).lower().strip()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    for ch in [".", ",", "-"]:
        texto = texto.replace(ch, " ")
    reemplazos = {
        "jiron": "jr", "jr": "jr",
        "avenida": "av", "avda": "av", "av": "av",
        "prolongacion": "prol", "prol": "prol",
        "pasaje": "psje", "psj": "psje", "psje": "psje",
        "calle": "cl", "cl": "cl",
    }
    palabras = texto.split()
    palabras = [reemplazos.get(p, p) for p in palabras]
    texto = " ".join(palabras)
    texto = re.sub(r"\bcuadra\b.*$", "", texto)
    texto = re.sub(r"\bcdra\b.*$", "", texto)
    texto = re.sub(r"[^a-z0-9\s]", " ", texto)
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()


def limpiar_espacios(texto: str) -> str:
    if not texto:
        return ""
    return str(texto).strip()


# ============================================================
# VALIDACION DE RUC
# ============================================================

def validar_ruc(ruc: str) -> bool:
    if not re.fullmatch(r"\d{11}", ruc):
        return False
    pesos = [5, 4, 3, 2, 7, 6, 5, 4, 3, 2]
    s = sum(int(ruc[i]) * pesos[i] for i in range(10))
    digito_verificador = (11 - (s % 11)) % 10
    return digito_verificador == int(ruc[10])


# ============================================================
# TARIFA
# ============================================================

def calcular_tarifa(distancia_ida_km: float) -> tuple:
    if distancia_ida_km <= 40:
        return Decimal("2.00"), Decimal("1.00")
    return Decimal("3.00"), Decimal("1.50")


# ============================================================
# LECTURA DEL EXCEL
# ============================================================

def leer_excel(ruta: str) -> tuple:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    rutas_ws = wb["tbl_rutas"]
    itin_ws = wb["tbl_itinerarios"]

    headers_rutas = [cel.value for cel in rutas_ws[1]]
    headers_itin = [cel.value for cel in itin_ws[1]]

    filas_rutas = []
    for fila in rutas_ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in fila):
            continue
        filas_rutas.append(dict(zip(headers_rutas, fila)))

    filas_itin = []
    for fila in itin_ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in fila):
            continue
        filas_itin.append(dict(zip(headers_itin, fila)))

    return filas_rutas, filas_itin


# ============================================================
# PROCESAMIENTO DE EMPRESAS
# ============================================================

def procesar_empresas(filas_rutas: list) -> dict:
    empresas_raw = {}
    for fila in filas_rutas:
        ruc = str(fila.get("ruc") or "").strip()
        empresa = limpiar_espacios(fila.get("empresa"))
        nombre_comercial = limpiar_espacios(fila.get("nombre comercial"))
        if not ruc:
            continue
        if ruc not in empresas_raw:
            empresas_raw[ruc] = {
                "razon_social": empresa,
                "nombre_comercial": nombre_comercial if nombre_comercial else None,
                "rutas": [],
            }
        else:
            if empresas_raw[ruc]["razon_social"] != empresa:
                print(f"  [ADVERTENCIA] RUC {ruc}: razon social inconsistente entre rutas")
                print(f"    '{empresas_raw[ruc]['razon_social']}' vs '{empresa}'")
            if nombre_comercial and empresas_raw[ruc]["nombre_comercial"] != nombre_comercial:
                print(f"  [ADVERTENCIA] RUC {ruc}: nombre comercial inconsistente entre rutas")
        empresas_raw[ruc]["rutas"].append(fila.get("id_ruta"))

    return empresas_raw


# ============================================================
# INSERCION EN BASE DE DATOS
# ============================================================

def importar(engine):
    print("\n=== Lectura del Excel ===")
    try:
        filas_rutas, filas_itin = leer_excel(EXCEL_PATH)
    except FileNotFoundError:
        print(f"ERROR: No se encontro el archivo: {EXCEL_PATH}")
        sys.exit(1)
    except Exception as e:
        print(f"ERROR al leer el Excel: {e}")
        sys.exit(1)

    print(f"  Rutas leidas: {len(filas_rutas)}")
    print(f"  Puntos leidos: {len(filas_itin)}")

    # Validar headers minimos
    headers_rutas = list(filas_rutas[0].keys()) if filas_rutas else []
    headers_itin = list(filas_itin[0].keys()) if filas_itin else []
    campos_requeridos_rutas = ["id_ruta", "codigo", "empresa", "ruc", "nombre comercial",
                                "origen", "destino", "distrito_origen", "distrito_destino",
                                "distancia_ida_km", "distancia_vuelta_km",
                                "tiempo_ida_min", "tiempo_vuelta_min",
                                "frecuencia_min", "horario"]
    campos_requeridos_itin = ["id_ruta", "sentido", "orden_secuencia", "via_calle", "cuadras", "distrito"]
    faltantes_rutas = [c for c in campos_requeridos_rutas if c not in headers_rutas]
    faltantes_itin = [c for c in campos_requeridos_itin if c not in headers_itin]
    if faltantes_rutas:
        print(f"ERROR: Columnas faltantes en tbl_rutas: {faltantes_rutas}")
        sys.exit(1)
    if faltantes_itin:
        print(f"ERROR: Columnas faltantes en tbl_itinerarios: {faltantes_itin}")
        sys.exit(1)

    print("\n=== Procesamiento de empresas ===")
    empresas_raw = procesar_empresas(filas_rutas)
    print(f"  Empresas detectadas: {len(empresas_raw)}")

    errores_ruc = []
    for ruc in empresas_raw:
        if not validar_ruc(ruc):
            errores_ruc.append(ruc)
            print(f"  [ERROR] RUC invalido: {ruc}")

    if errores_ruc:
        print(f"\nERROR: {len(errores_ruc)} RUC(s) invalido(s). Corrige el Excel antes de importar.")
        sys.exit(1)

    # Agrupar puntos por ruta y sentido
    puntos_por_ruta = {}
    for fila in filas_itin:
        id_ruta = fila.get("id_ruta")
        sentido = fila.get("sentido")
        if not id_ruta or not sentido:
            continue
        clave = (id_ruta, sentido)
        if clave not in puntos_por_ruta:
            puntos_por_ruta[clave] = []
        puntos_por_ruta[clave].append(fila)

    # Ordenar puntos por orden_secuencia
    for clave in puntos_por_ruta:
        puntos_por_ruta[clave].sort(key=lambda x: int(x.get("orden_secuencia") or 0))

    print("\n=== Insercion en base de datos ===")

    with engine.begin() as conn:
        # Limpiar tablas existentes
        conn.execute(text("DELETE FROM puntos_recorrido"))
        conn.execute(text("DELETE FROM sentidos"))
        conn.execute(text("DELETE FROM rutas"))
        conn.execute(text("DELETE FROM empresas"))
        print("  Tablas limpiadas")

        empresa_ids = {}
        for ruc, data in empresas_raw.items():
            result = conn.execute(
                text("""
                    INSERT INTO empresas (nombre_comercial, razon_social, ruc)
                    VALUES (:nc, :rs, :ruc)
                """),
                {"nc": data["nombre_comercial"], "rs": data["razon_social"], "ruc": ruc}
            )
            empresa_ids[ruc] = result.lastrowid

        print(f"  Empresas insertadas: {len(empresa_ids)}")

        ruta_ids = {}
        errores_ruta = 0

        for fila in filas_rutas:
            id_ruta = str(fila.get("id_ruta") or "").strip()
            nombre = limpiar_espacios(fila.get("codigo"))
            ruc = str(fila.get("ruc") or "").strip()
            empresa_id = empresa_ids.get(ruc)

            origen = limpiar_espacios(fila.get("origen"))
            destino = limpiar_espacios(fila.get("destino"))
            frecuencia = int(fila.get("frecuencia_min") or 0) or None
            horario = limpiar_espacios(fila.get("horario") or "06:00 - 20:00")

            try:
                hi_str, hf_str = horario.split("-")
                hi = hi_str.strip()
                hf = hf_str.strip()
                hi_fmt = hi if ":" in hi else f"{hi}:00"
                hf_fmt = hf if ":" in hf else f"{hf}:00"
            except Exception:
                hi_fmt = "06:00"
                hf_fmt = "20:00"

            flota = int(fila.get("flota_maxima") or 0) or None
            reten = int(fila.get("reten") or 0) or None
            velocidad = None
            try:
                velocidad = Decimal(str(fila.get("velocidad_kmh") or ""))
            except (InvalidOperation, ValueError):
                pass
            pasajeros = int(fila.get("pasajeros_vuelta") or 0) or None
            ipk = None
            try:
                ipk = Decimal(str(fila.get("ipk") or ""))
            except (InvalidOperation, ValueError):
                pass

            distancia_ida = 0.0
            try:
                distancia_ida = float(fila.get("distancia_ida_km") or 0)
            except (ValueError, TypeError):
                pass

            tarifa_gen, tarifa_med = calcular_tarifa(distancia_ida)

            try:
                result = conn.execute(
                    text("""
                        INSERT INTO rutas (
                            codigo, nombre, empresa_id, frecuencia_general_min,
                            horario_inicio, horario_fin,
                            flota_maxima, reten_unidades, velocidad_promedio_kmh,
                            pasajeros_por_vuelta, ipk,
                            tarifa_general, tarifa_medio_pasaje, tarifa_es_provisional,
                            fuente, activo
                        )
                        VALUES (
                            :cod, :nom, :eid, :freq,
                            :hi, :hf,
                            :flota, :reten, :vel,
                            :pasj, :ipk,
                            :tg, :tm, FALSE,
                            'EXCEL_V4', TRUE
                        )
                    """),
                    {
                        "cod": id_ruta, "nom": nombre, "eid": empresa_id,
                        "freq": frecuencia,
                        "hi": hi_fmt, "hf": hf_fmt,
                        "flota": flota, "reten": reten, "vel": velocidad,
                        "pasj": pasajeros, "ipk": ipk,
                        "tg": tarifa_gen, "tm": tarifa_med,
                    }
                )
                ruta_ids[id_ruta] = result.lastrowid
            except IntegrityError:
                print(f"  [ERROR] Ruta duplicada: {id_ruta}")
                errores_ruta += 1

        print(f"  Rutas insertadas: {len(ruta_ids)} (errores: {errores_ruta})")

        # Sentidos
        sentidos_insertados = 0
        errores_sentido = 0
        for fila in filas_rutas:
            id_ruta = str(fila.get("id_ruta") or "").strip()
            ruta_id = ruta_ids.get(id_ruta)
            if not ruta_id:
                continue

            distancia_ida = 0.0
            distancia_vuelta = 0.0
            tiempo_ida = 0
            tiempo_vuelta = 0
            try:
                distancia_ida = float(fila.get("distancia_ida_km") or 0)
            except (ValueError, TypeError):
                pass
            try:
                distancia_vuelta = float(fila.get("distancia_vuelta_km") or 0)
            except (ValueError, TypeError):
                pass
            try:
                tiempo_ida = int(fila.get("tiempo_ida_min") or 0)
            except (ValueError, TypeError):
                pass
            try:
                tiempo_vuelta = int(fila.get("tiempo_vuelta_min") or 0)
            except (ValueError, TypeError):
                pass

            origen_ida = limpiar_espacios(fila.get("origen"))
            destino_ida = limpiar_espacios(fila.get("destino"))
            origen_vuelta = limpiar_espacios(fila.get("destino"))
            destino_vuelta = limpiar_espacios(fila.get("origen"))

            for tipo, ori, des, dist, tiempo in [
                ("IDA", origen_ida, destino_ida, distancia_ida, tiempo_ida),
                ("VUELTA", origen_vuelta, destino_vuelta, distancia_vuelta, tiempo_vuelta),
            ]:
                try:
                    result = conn.execute(
                        text("""
                            INSERT INTO sentidos (
                                ruta_id, tipo, origen, destino,
                                distancia_km, tiempo_total_min
                            )
                            VALUES (:rid, :tipo, :ori, :des, :dist, :tiempo)
                        """),
                        {"rid": ruta_id, "tipo": tipo, "ori": ori, "des": des,
                         "dist": dist, "tiempo": tiempo}
                    )
                    sentidos_insertados += 1
                except IntegrityError:
                    print(f"  [ERROR] Sentido duplicado: {id_ruta} {tipo}")
                    errores_sentido += 1

        print(f"  Sentidos insertados: {sentidos_insertados} (errores: {errores_sentido})")

        # Puntos de recorrido
        puntos_insertados = 0
        errores_punto = 0
        sin_sentido = 0

        for (id_ruta, sentido), puntos in puntos_por_ruta.items():
            ruta_id = ruta_ids.get(id_ruta)
            if not ruta_id:
                sin_sentido += 1
                continue

            sentido_row = conn.execute(
                text("SELECT id FROM sentidos WHERE ruta_id = :rid AND tipo = :tipo"),
                {"rid": ruta_id, "tipo": sentido}
            ).fetchone()

            if not sentido_row:
                print(f"  [ERROR] No se encontro sentido {id_ruta} {sentido}")
                errores_sentido += 1
                continue

            sentido_id = sentido_row[0]

            for idx, p in enumerate(puntos, start=1):
                orden = int(p.get("orden_secuencia") or idx)
                nombre_orig = limpiar_espacios(p.get("via_calle"))
                nombre_norm = normalizar(nombre_orig)
                cuadra = limpiar_espacios(p.get("cuadras"))
                distrito = limpiar_espacios(p.get("distrito"))

                if not nombre_orig:
                    continue

                try:
                    conn.execute(
                        text("""
                            INSERT INTO puntos_recorrido (
                                sentido_id, orden, nombre_original, nombre_normalizado,
                                tipo_via, cuadra, distrito
                            )
                            VALUES (:sid, :ord, :no, :nn, :tv, :cd, :dis)
                        """),
                        {
                            "sid": sentido_id, "ord": orden,
                            "no": nombre_orig, "nn": nombre_norm,
                            "tv": None, "cd": cuadra, "dis": distrito,
                        }
                    )
                    puntos_insertados += 1
                except IntegrityError:
                    print(f"  [ERROR] Punto duplicado: {id_ruta} {sentido} orden {orden}")
                    errores_punto += 1

        print(f"  Puntos insertados: {puntos_insertados} (errores: {errores_punto})")
        if sin_sentido:
            print(f"  Rutas sin sentido (no encontradas): {sin_sentido}")

    print("\n=== Resumen ===")
    print(f"  Empresas procesadas: {len(empresa_ids)}")
    print(f"  Rutas insertadas: {len(ruta_ids)}")
    print(f"  Sentidos insertados: {sentidos_insertados}")
    print(f"  Puntos insertados: {puntos_insertados}")
    advertencias = errores_ruta + errores_sentido + errores_punto
    print(f"  Errores: {advertencias}")
    if advertencias == 0:
        print("  Importacion completada correctamente")
    else:
        print("  Revisa los errores antes de continuar")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        EXCEL_PATH = sys.argv[1]
    print("Importador Excel v4 -> MySQL")
    print(f"Base de datos: {DB_NAME}@{DB_HOST}:{DB_PORT}")
    print(f"Archivo: {EXCEL_PATH}")

    engine = create_engine(DATABASE_URL, pool_pre_ping=True)
    importar(engine)
